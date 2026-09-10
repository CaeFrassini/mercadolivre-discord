import asyncio
import base64
import hashlib
import tempfile
import time
import unittest
from datetime import datetime, timedelta
from decimal import Decimal
from urllib.parse import parse_qs, urlsplit
from unittest.mock import AsyncMock
from openpyxl import load_workbook
from aiohttp.test_utils import TestClient, TestServer
from main import MercadoLivre, Store, SafeError, UTC, paid_totals, workbook, web_app, Bot


class Tests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.store = Store(self.temp.name)
        self.ml = MercadoLivre(None, self.store, '123', 'fake', 'https://example.com/oauth/callback')
        self.store.set('tokens', {'user_id': 5})

    def tearDown(self):
        self.store.db.close()
        self.temp.cleanup()

    def test_pkce_state_and_replay(self):
        url = self.ml.authorization_link()
        p = parse_qs(urlsplit(url).query)
        with self.assertRaises(SafeError):
            self.ml.consume_state('bad')
        verifier = self.ml.consume_state(p['state'][0])
        expected = base64.urlsafe_b64encode(hashlib.sha256(verifier.encode()).digest()).rstrip(b'=').decode()
        self.assertEqual(p['code_challenge'][0], expected)
        with self.assertRaises(SafeError):
            self.ml.consume_state(p['state'][0])

    def test_expired_and_persistent(self):
        p = parse_qs(urlsplit(self.ml.authorization_link()).query)
        pending = self.store.get('oauth_pending')
        pending['expires'] = time.time() - 1
        self.store.set('oauth_pending', pending)
        with self.assertRaises(SafeError):
            self.ml.consume_state(p['state'][0])
        second = Store(self.temp.name)
        self.assertEqual(second.get('tokens')['user_id'], 5)
        second.db.close()

    async def test_orders_boundary_dedup_and_pages(self):
        start = datetime(2026, 9, 10, tzinfo=UTC)
        def order(i, when):
            return {'id': i, 'date_created': when.isoformat()}
        self.ml.get = AsyncMock(side_effect=[
            {'paging': {'total': 4}, 'results': [order(1, start), order(2, start + timedelta(minutes=30))]},
            {'paging': {'total': 4}, 'results': [order(2, start + timedelta(minutes=30)), order(3, start + timedelta(hours=1))]}])
        rows = await self.ml.orders(start, start + timedelta(hours=1))
        self.assertEqual([o['id'] for o in rows], [1, 2])
        self.assertEqual(self.ml.get.await_count, 2)

    async def test_items_scan_and_failure(self):
        self.ml.get = AsyncMock(side_effect=[{'results': ['MLB1'], 'scroll_id': 'a'}, {'results': []},
                                             [{'code': 200, 'body': {'id': 'MLB1', 'status': 'active'}}]])
        self.assertEqual((await self.ml.active_items())[0]['id'], 'MLB1')
        self.ml.get = AsyncMock(side_effect=[{'results': ['MLB1']}, [{'code': 403}]])
        with self.assertRaises(SafeError):
            await self.ml.active_items()

    async def test_callback_invalid_state_never_exchanges(self):
        self.ml.token_request = AsyncMock()
        async with TestClient(TestServer(web_app(self.ml))) as client:
            r = await client.get('/oauth/callback?code=fake&state=wrong')
            self.assertEqual(r.status, 400)
            self.ml.token_request.assert_not_awaited()
            r = await client.get('/health')
            self.assertEqual(r.status, 200)

    async def test_callback_pkce_success(self):
        p = parse_qs(urlsplit(self.ml.authorization_link()).query)
        self.ml.token_request = AsyncMock()
        async with TestClient(TestServer(web_app(self.ml))) as client:
            r = await client.get('/oauth/callback', params={'code': 'fake', 'state': p['state'][0]})
            self.assertEqual(r.status, 200)
            self.assertIn('code_verifier', self.ml.token_request.call_args.args[0])
            self.assertEqual(r.headers['Cache-Control'], 'no-store')

    async def test_refresh_single_flight(self):
        self.store.set('tokens', {'user_id': 5, 'access_token': 'old', 'refresh_token': 'refresh', 'expires_at': 0})
        async def exchange(fields):
            await asyncio.sleep(0)
            data = {'user_id': 5, 'access_token': 'new', 'refresh_token': 'rotated', 'expires_at': time.time() + 3600}
            self.store.set('tokens', data)
            return data
        self.ml.token_request = AsyncMock(side_effect=exchange)
        result = await asyncio.gather(self.ml.access(), self.ml.access())
        self.assertEqual(result, ['new', 'new'])
        self.ml.token_request.assert_awaited_once()
        self.assertEqual(self.store.get('tokens')['refresh_token'], 'rotated')

    async def test_commands_and_owner_restriction(self):
        from types import SimpleNamespace
        bot = Bot(self.ml, self.store, owner_id=7, channel_id=9)
        bot.channel = SimpleNamespace(guild=SimpleNamespace(id=10))
        self.assertEqual({c.name for c in bot.tree.get_commands()}, {'conectar', 'status', 'relatorio'})
        interaction = SimpleNamespace(user=SimpleNamespace(id=8), guild_id=10,
                                      response=SimpleNamespace(send_message=AsyncMock()))
        self.assertFalse(await bot.permitted(interaction))
        self.assertTrue(await bot.permitted(interaction, owner_only=False))
        interaction.guild_id = 11
        self.assertFalse(await bot.permitted(interaction, owner_only=False))
        interaction.guild_id = None
        self.assertFalse(await bot.permitted(interaction, owner_only=False))
        interaction.guild_id = 10
        interaction.user.id = 7
        self.assertTrue(await bot.permitted(interaction))
        await bot.close()

    def test_financial_status_currency_and_excel(self):
        rows = [dict(id=1, status='paid', total_amount=0.1, currency_id='BRL'),
                dict(id=2, status='paid', total_amount=0.2, currency_id='BRL'),
                dict(id=3, status='cancelled', total_amount=999, currency_id='BRL'),
                dict(id=4, status='paid', total_amount=2, currency_id='USD')]
        count, totals = paid_totals(rows)
        self.assertEqual(count, 3)
        self.assertEqual(totals, {'BRL': Decimal('0.3'), 'USD': Decimal('2')})
        now = datetime.now(UTC)
        output = workbook([{'id': 'MLB1', 'title': '=1+1', 'status': 'active'}], rows, now - timedelta(hours=1), now)
        wb = load_workbook(output)
        self.assertEqual(len(wb.sheetnames), 6)
        self.assertEqual(wb['Anuncios ativos']['B2'].data_type, 's')
        self.assertEqual(wb['Pedidos do periodo'].max_row, 5)
        self.assertIsNone(wb['Anuncios ativos']['E2'].value)


if __name__ == '__main__':
    unittest.main()
