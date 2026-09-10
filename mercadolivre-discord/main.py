"""Relatórios privados Mercado Livre → Discord. Python 3.12+."""
import asyncio
import base64
import hashlib
import io
import json
import logging
import os
import secrets
import sqlite3
import time
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from urllib.parse import urlencode
from zoneinfo import ZoneInfo

import aiohttp
from aiohttp import web
import discord
from discord import app_commands
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

UTC = timezone.utc
BR = ZoneInfo('America/Sao_Paulo')
API = 'https://api.mercadolibre.com'
log = logging.getLogger('relatorios')


class SafeError(Exception):
    """Somente mensagens próprias, sem corpos HTTP, URLs ou credenciais."""


class Store:
    def __init__(self, folder):
        Path(folder).mkdir(parents=True, exist_ok=True)
        path = Path(folder) / 'bot.sqlite3'
        self.db = sqlite3.connect(path)
        os.chmod(path, 0o600)
        self.db.execute('CREATE TABLE IF NOT EXISTS kv (k TEXT PRIMARY KEY, v TEXT NOT NULL)')
        self.db.commit()

    def get(self, key, default=None):
        row = self.db.execute('SELECT v FROM kv WHERE k=?', (key,)).fetchone()
        return json.loads(row[0]) if row else default

    def set(self, key, value):
        with self.db:
            self.db.execute('INSERT OR REPLACE INTO kv VALUES (?, ?)', (key, json.dumps(value)))

    def pop(self, key):
        value = self.get(key)
        with self.db:
            self.db.execute('DELETE FROM kv WHERE k=?', (key,))
        return value


def stamp(value):
    return value.astimezone(UTC).isoformat(timespec='milliseconds').replace('+00:00', 'Z')


def parsed(value):
    return datetime.fromisoformat(value.replace('Z', '+00:00'))


class MercadoLivre:
    def __init__(self, session, store, client_id, secret, redirect):
        self.session, self.store = session, store
        self.client_id, self.secret, self.redirect = client_id, secret, redirect
        self.lock = asyncio.Lock()

    def authorization_link(self):
        state, verifier = secrets.token_urlsafe(32), secrets.token_urlsafe(64)
        challenge = base64.urlsafe_b64encode(hashlib.sha256(verifier.encode()).digest()).rstrip(b'=').decode()
        self.store.set('oauth_pending', {'hash': hashlib.sha256(state.encode()).hexdigest(),
                                       'verifier': verifier, 'expires': time.time() + 600})
        return 'https://auth.mercadolivre.com.br/authorization?' + urlencode({
            'response_type': 'code', 'client_id': self.client_id, 'redirect_uri': self.redirect,
            'state': state, 'code_challenge': challenge, 'code_challenge_method': 'S256'})

    def consume_state(self, state):
        pending = self.store.get('oauth_pending')
        digest = hashlib.sha256(state.encode()).hexdigest()
        if not pending or pending['expires'] < time.time() or not secrets.compare_digest(pending['hash'], digest):
            raise SafeError('Link inválido ou expirado. Use /conectar novamente.')
        self.store.pop('oauth_pending')
        return pending['verifier']

    async def token_request(self, fields):
        # Não repetir POST automaticamente: códigos e refresh tokens podem ser de uso único.
        async with self.session.post(API + '/oauth/token', data={
            'client_id': self.client_id, 'client_secret': self.secret, **fields}) as response:
            if response.status != 200:
                raise SafeError(f'Autorização falhou (HTTP {response.status}). Use /conectar novamente.')
            data = await response.json()
        if not all(data.get(k) for k in ('access_token', 'refresh_token', 'user_id', 'expires_in')):
            raise SafeError('Resposta de autorização incompleta. Use /conectar novamente.')
        old = self.store.get('tokens')
        if old and str(old['user_id']) != str(data['user_id']):
            raise SafeError('Conta diferente da já conectada. Use a mesma conta vendedora.')
        data['expires_at'] = time.time() + int(data['expires_in'])
        self.store.set('tokens', data)  # Transação: salva também o novo refresh token.
        return data

    async def authorize(self, code, state):
        verifier = self.consume_state(state)
        async with self.lock:
            await self.token_request({'grant_type': 'authorization_code', 'code': code,
                                      'redirect_uri': self.redirect, 'code_verifier': verifier})

    async def access(self, rejected=None):
        async with self.lock:
            data = self.store.get('tokens')
            if not data:
                raise SafeError('Conta não conectada. Use /conectar.')
            if data['expires_at'] < time.time() + 120 or data['access_token'] == rejected:
                data = await self.token_request({'grant_type': 'refresh_token', 'refresh_token': data['refresh_token']})
            return data['access_token']

    async def get(self, path, params=None):
        token = await self.access()
        for attempt in range(5):
            try:
                async with self.session.get(API + path, params=params, headers={'Authorization': 'Bearer ' + token}) as r:
                    if r.status == 200:
                        return await r.json()
                    status = r.status
                    try:
                        delay = min(30, max(1, float(r.headers.get('Retry-After', 2 ** attempt))))
                    except ValueError:
                        delay = 2 ** attempt
                if status == 401 and attempt < 4:
                    token = await self.access(rejected=token)
                    continue
                if status == 429 or status >= 500:
                    await asyncio.sleep(delay)
                    continue
                raise SafeError(f'Consulta ao Mercado Livre falhou (HTTP {status}). Confira permissões e autorização.')
            except (aiohttp.ClientError, asyncio.TimeoutError):
                if attempt == 4:
                    raise SafeError('Mercado Livre indisponível. Tente novamente mais tarde.') from None
                await asyncio.sleep(2 ** attempt)
        raise SafeError('Mercado Livre limitou ou não concluiu a consulta. Tente novamente.')

    @property
    def seller(self):
        data = self.store.get('tokens')
        if not data:
            raise SafeError('Use /conectar primeiro.')
        return data['user_id']

    async def active_items(self):
        ids, seen, scroll = [], set(), None
        while True:
            params = {'status': 'active', 'search_type': 'scan', 'limit': 100}
            if scroll:
                params['scroll_id'] = scroll
            data = await self.get(f'/users/{self.seller}/items/search', params)
            page = data.get('results')
            if not isinstance(page, list):
                raise SafeError('Resposta de anúncios inesperada; relatório não enviado.')
            if not page:
                break
            fresh = [x for x in page if x not in seen]
            if not fresh:
                raise SafeError('Paginação de anúncios repetida; relatório não enviado como completo.')
            ids.extend(fresh)
            seen.update(fresh)
            scroll = data.get('scroll_id')
            if not scroll:
                if len(page) < 100:
                    break
                raise SafeError('Paginação de anúncios incompleta; relatório interrompido.')
        items = []
        for i in range(0, len(ids), 20):
            batch_ids = ids[i:i+20]
            batch = await self.get('/items', {'ids': ','.join(batch_ids), 'include_attributes': 'all'})
            if not isinstance(batch, list) or len(batch) != len(batch_ids):
                raise SafeError('Detalhes de anúncios incompletos; tente novamente.')
            for entry in batch:
                if entry.get('code') != 200:
                    raise SafeError('Um anúncio não pôde ser consultado; relatório interrompido.')
                body = entry['body']
                # O anúncio pode ter sido pausado entre a busca e a leitura.
                if body.get('status') == 'active':
                    items.append(body)
        return items

    async def orders(self, start, end):
        if start >= end:
            return []
        params = {'seller': self.seller, 'order.date_created.from': stamp(start),
                  'order.date_created.to': stamp(end), 'sort': 'date_asc', 'limit': 50, 'offset': 0}
        data = await self.get('/orders/search', params)
        total = data.get('paging', {}).get('total')
        if not isinstance(total, int) or not isinstance(data.get('results'), list):
            raise SafeError('Resposta de vendas inesperada; relatório interrompido.')
        # Divide o intervalo antes de ultrapassar janelas de paginação usuais.
        if total > 1000:
            if end - start <= timedelta(seconds=1):
                raise SafeError('Muitas vendas no mesmo segundo; relatório requer paginação especializada.')
            midpoint = start + (end - start) / 2
            rows = await self.orders(start, midpoint) + await self.orders(midpoint, end)
        else:
            rows = list(data['results'])
            offset = len(rows)
            while offset < total:
                if not data['results']:
                    raise SafeError('Paginação de vendas incompleta; relatório interrompido.')
                params['offset'] = offset
                data = await self.get('/orders/search', params)
                if not isinstance(data.get('results'), list):
                    raise SafeError('Resposta de vendas incompleta.')
                rows.extend(data['results'])
                offset += len(data['results'])
        # Janela [início, fim), sem contar pedidos duas vezes no limite.
        return list({str(o['id']): o for o in rows if start <= parsed(o['date_created']) < end}.values())


def paid_totals(orders):
    totals = defaultdict(Decimal)
    count = 0
    for order in orders:
        if order.get('status') == 'paid':
            totals[order.get('currency_id', 'N/D')] += Decimal(str(order.get('total_amount') or 0))
            count += 1
    return count, dict(totals)


def money(totals):
    return ' | '.join(f'{currency} {value:,.2f}' for currency, value in sorted(totals.items())) or '0,00'


def safe_cell(value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub('', value)[:32000]
        if value.lstrip().startswith(('=', '+', '-', '@')):
            value = "'" + value
    return value


def workbook(items, orders, start, end):
    wb = Workbook()
    wb.remove(wb.active)
    def sheet(name, headers, rows):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append([safe_cell(v) for v in row])
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='17365D')
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(55, max(15, len(str(col[0].value)) + 3))
        return ws
    sheet('Anuncios ativos', ['ID', 'Título', 'Preço atual', 'Moeda', 'Estoque informado', 'Vendidos acumulados API',
          'Status', 'Tipo', 'Condição', 'Categoria', 'SKU', 'Link', 'Criado em', 'Atualizado em', 'Frete grátis', 'Logística'],
          ([x.get('id'), x.get('title'), x.get('price'), x.get('currency_id'), x.get('available_quantity'),
            x.get('sold_quantity'), x.get('status'), x.get('listing_type_id'), x.get('condition'), x.get('category_id'),
            x.get('seller_custom_field') or next((a.get('value_name') for a in x.get('attributes', []) if a.get('id') == 'SELLER_SKU'), None),
            x.get('permalink'), x.get('date_created'), x.get('last_updated'), x.get('shipping', {}).get('free_shipping'),
            x.get('shipping', {}).get('logistic_type')] for x in items))
    sheet('Variacoes', ['Anúncio', 'Variação', 'Preço', 'Estoque informado', 'Vendidos API', 'SKU', 'Atributos'],
          ([x['id'], str(v.get('id', '')), v.get('price'), v.get('available_quantity'), v.get('sold_quantity'),
            v.get('seller_custom_field'), v.get('attribute_combinations', [])] for x in items for v in x.get('variations', [])))
    sheet('Pedidos do periodo', ['Pedido', 'Criado em', 'Status', 'Moeda', 'Total pedido (bruto)', 'Pago informado', 'Envio ID', 'Pack ID'],
          ([str(o['id']), o.get('date_created'), o.get('status'), o.get('currency_id'), o.get('total_amount'), o.get('paid_amount'),
            str((o.get('shipping') or {}).get('id') or ''), str(o.get('pack_id') or '')] for o in orders))
    sheet('Itens vendidos', ['Pedido', 'Anúncio', 'Título', 'Variação', 'Quantidade', 'Preço unitário', 'Taxa unitária informada', 'Moeda'],
          ([str(o['id']), v.get('item', {}).get('id'), v.get('item', {}).get('title'), str(v.get('item', {}).get('variation_id') or ''),
            v.get('quantity'), v.get('unit_price'), v.get('sale_fee'), v.get('currency_id', o.get('currency_id'))]
           for o in orders for v in o.get('order_items', [])))
    sheet('Pagamentos', ['Pedido', 'Pagamento ID', 'Status', 'Detalhe', 'Valor transação', 'Valor reembolsado', 'Moeda'],
          ([str(o['id']), str(p.get('id', '')), p.get('status'), p.get('status_detail'), p.get('transaction_amount'),
            p.get('transaction_amount_refunded'), p.get('currency_id', o.get('currency_id'))] for o in orders for p in o.get('payments', [])))
    sheet('Notas', ['Campo', 'Explicação'], [
        ['Início inclusivo', start.astimezone(BR).isoformat()], ['Fim exclusivo', end.astimezone(BR).isoformat()],
        ['Pedidos', 'Todos os status; filtro pela criação, não pela data de pagamento. Inclui itens hoje pausados.'],
        ['Resumo', 'Soma total_amount apenas de pedidos atualmente paid. Valor bruto, não lucro ou saldo disponível.'],
        ['Estoque', 'Valor devolvido pela API; pode diferir do painel em modelos de estoque compartilhado. Não some variações com o total do anúncio.'],
        ['Taxas', 'sale_fee unitária devolvida pela API, quando disponível. Não representa todos os custos.'],
        ['Limites', 'Não inclui conciliação financeira, custo do fornecedor, tarifas completas, status detalhado de transporte, descrição completa ou dados pessoais de compradores.'],
        ['Campos vazios', 'Dado não devolvido pela API; não significa zero.'],
        ['Atualização', 'Pedidos e anúncios são lidos em sequência; a conta pode mudar durante a consulta.'],
        ['Histórico', 'Mudanças tardias de pagamento/cancelamento não reescrevem mensagens antigas. Gere /relatorio para consultar novamente.']])
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


class Bot(discord.Client):
    def __init__(self, ml, store, owner_id, channel_id):
        super().__init__(intents=discord.Intents.none(), allowed_mentions=discord.AllowedMentions.none())
        self.ml, self.store = ml, store
        self.owner_id, self.channel_id = owner_id, channel_id
        self.tree = app_commands.CommandTree(self)
        self.report_lock = asyncio.Lock()
        self.worker = None
        self.channel = None
        self.register_commands()

    async def permitted(self, interaction, *, owner_only=True):
        if interaction.guild_id != self.channel.guild.id:
            await interaction.response.send_message('Use este comando no servidor configurado.', ephemeral=True)
            return False
        if owner_only and interaction.user.id != self.owner_id:
            await interaction.response.send_message('Comando restrito ao proprietário configurado.', ephemeral=True)
            return False
        return True

    def register_commands(self):
        @self.tree.command(name='conectar', description='Autorizar sua conta do Mercado Livre (somente proprietário).')
        async def conectar(interaction: discord.Interaction):
            if not await self.permitted(interaction):
                return
            url = self.ml.authorization_link()
            view = discord.ui.View(timeout=600)
            view.add_item(discord.ui.Button(label='Autorizar Mercado Livre', url=url))
            await interaction.response.send_message('Abra o link em até 10 minutos e autorize a conta vendedora. Não compartilhe este link.', view=view, ephemeral=True)

        @self.tree.command(name='status', description='Verificar se o Mercado Livre está conectado.')
        async def status(interaction: discord.Interaction):
            if not await self.permitted(interaction):
                return
            text = 'Conta conectada. Relatórios automáticos a cada hora.' if self.store.get('tokens') else 'Use /conectar para autorizar o Mercado Livre.'
            await interaction.response.send_message(text, ephemeral=True)

        @self.tree.command(name='relatorio', description='Gerar planilha dos anúncios ativos e pedidos recentes.')
        @app_commands.describe(dias='Período de pedidos: entre 1 e 30 dias.')
        async def relatorio(interaction: discord.Interaction, dias: app_commands.Range[int, 1, 30] = 1):
            if not await self.permitted(interaction, owner_only=False):
                return
            if self.report_lock.locked():
                await interaction.response.send_message('Já existe um relatório em andamento. Aguarde.', ephemeral=True)
                return
            await interaction.response.defer(ephemeral=True, thinking=True)
            try:
                end = datetime.now(UTC)
                async with self.report_lock:
                    content, file = await self.make_report(end - timedelta(days=dias), end)
                await interaction.followup.send(content, file=file, ephemeral=True)
            except Exception as exc:
                await interaction.followup.send(self.error_text(exc), ephemeral=True)

    @staticmethod
    def error_text(exc):
        log.warning('Falha na operação: %s', type(exc).__name__)
        return str(exc) if isinstance(exc, SafeError) else 'Falha ao gerar/enviar relatório. Verifique a conexão e as permissões do bot e tente novamente.'

    async def setup_hook(self):
        self.channel = await self.fetch_channel(self.channel_id)
        if not isinstance(self.channel, discord.TextChannel):
            raise SafeError('DISCORD_CHANNEL_ID deve ser de um canal de texto de servidor.')
        guild = discord.Object(id=self.channel.guild.id)
        self.tree.copy_global_to(guild=guild)
        await self.tree.sync(guild=guild)
        self.worker = asyncio.create_task(self.hourly())
        log.info('Bot iniciado; comandos disponíveis no servidor configurado.')

    async def make_report(self, start, end):
        items = await self.ml.active_items()
        orders = await self.ml.orders(start, end)
        count, totals = paid_totals(orders)
        low = sum(1 for i in items if i.get('available_quantity') is not None and i['available_quantity'] <= 5)
        blob = await asyncio.to_thread(workbook, items, orders, start, end)
        if blob.getbuffer().nbytes > min(self.channel.guild.filesize_limit, 9_000_000):
            raise SafeError('Planilha excede o limite de envio. Solicite um período menor ou adapte o envio em partes.')
        content = (f'**Mercado Livre — relatório**\n'
                   f'{start.astimezone(BR):%d/%m %H:%M} até {end.astimezone(BR):%d/%m %H:%M} (Brasília)\n'
                   f'Anúncios ativos agora: **{len(items)}** | Estoque informado ≤ 5: **{low}**\n'
                   f'Pedidos criados no período: **{len(orders)}** | Atualmente pagos: **{count}**\n'
                   f'Total bruto dos pedidos pagos: **{money(totals)}**\n'
                   'Valores brutos, sem conciliação de taxas, frete e reembolsos. Detalhes e limites na planilha.')
        return content, discord.File(blob, filename=f'relatorio_{end.astimezone(BR):%Y%m%d_%H%M}.xlsx')

    async def hourly(self):
        await self.wait_until_ready()
        # Primeira execução após a próxima virada da hora; reinícios não enviam imediatamente.
        next_end = datetime.now(UTC).replace(minute=0, second=0, microsecond=0) + timedelta(hours=1)
        while not self.is_closed():
            await asyncio.sleep(max(0, (next_end - datetime.now(UTC)).total_seconds()))
            end = next_end
            next_end = datetime.now(UTC).replace(minute=0, second=0, microsecond=0) + timedelta(hours=1)
            if not self.store.get('tokens') or self.store.get('last_sent') == stamp(end):
                continue
            try:
                async with self.report_lock:
                    content, file = await self.make_report(end - timedelta(hours=1), end)
                    # Acumulado do dia é calculado separadamente até o fim da janela.
                    local_day = (end - timedelta(microseconds=1)).astimezone(BR).replace(hour=0, minute=0, second=0, microsecond=0)
                    day_orders = await self.ml.orders(local_day, end)
                    n, amounts = paid_totals(day_orders)
                    content += f'\nAcumulado de {local_day:%d/%m} até esse horário: **{n} pedidos pagos | {money(amounts)}**.'
                    await self.channel.send(content, file=file)
                    self.store.set('last_sent', stamp(end))
            except Exception as exc:
                message = self.error_text(exc)
                try:
                    await self.channel.send('⚠️ Relatório da hora não concluído. ' + message)
                except Exception:
                    log.warning('Não foi possível enviar aviso ao canal.')

    async def close(self):
        if self.worker:
            self.worker.cancel()
            await asyncio.gather(self.worker, return_exceptions=True)
        await super().close()


def web_app(ml):
    app = web.Application()
    async def health(request):
        return web.Response(text='Serviço ativo. Use /conectar no Discord para autorizar.')
    async def callback(request):
        headers = {'Cache-Control': 'no-store', 'Referrer-Policy': 'no-referrer', 'X-Content-Type-Options': 'nosniff'}
        try:
            state, code = request.query.get('state', ''), request.query.get('code', '')
            if request.query.get('error'):
                ml.consume_state(state)
                raise SafeError('Autorização não concedida. Volte ao Discord e use /conectar.')
            if not code or not state:
                raise SafeError('Retorno incompleto. Inicie com /conectar no Discord.')
            await ml.authorize(code, state)
            return web.Response(text='Mercado Livre conectado! Volte ao Discord e use /relatorio para testar.', headers=headers)
        except Exception as exc:
            text = str(exc) if isinstance(exc, SafeError) else 'Falha na conexão. Gere um novo link com /conectar.'
            return web.Response(text=text, status=400, headers=headers)
    app.router.add_get('/', health)
    app.router.add_get('/health', health)
    app.router.add_get('/oauth/callback', callback)
    return app


async def main():
    names = ['ML_CLIENT_ID', 'ML_CLIENT_SECRET', 'ML_REDIRECT_URI', 'DISCORD_TOKEN', 'DISCORD_CHANNEL_ID', 'DISCORD_OWNER_ID']
    missing = [name for name in names if not os.environ.get(name)]
    if missing:
        raise SafeError('Variáveis ausentes: ' + ', '.join(missing))
    cfg = {key: os.environ[key].strip() for key in names}
    if not cfg['ML_REDIRECT_URI'].startswith('https://') or not cfg['ML_REDIRECT_URI'].endswith('/oauth/callback'):
        raise SafeError('ML_REDIRECT_URI deve ser HTTPS e terminar em /oauth/callback.')
    folder = os.getenv('DATA_DIR', '/data')
    store = Store(folder)
    async with aiohttp.ClientSession(timeout=aiohttp.ClientTimeout(total=45)) as session:
        ml = MercadoLivre(session, store, cfg['ML_CLIENT_ID'], cfg['ML_CLIENT_SECRET'], cfg['ML_REDIRECT_URI'])
        bot = Bot(ml, store, int(cfg['DISCORD_OWNER_ID']), int(cfg['DISCORD_CHANNEL_ID']))
        runner = web.AppRunner(web_app(ml), access_log=None)  # Não registrar códigos OAuth nas URLs.
        await runner.setup()
        await web.TCPSite(runner, '0.0.0.0', int(os.getenv('PORT', '8080'))).start()
        try:
            async with bot:
                await bot.start(cfg['DISCORD_TOKEN'])
        finally:
            await runner.cleanup()
            store.db.close()


if __name__ == '__main__':
    os.umask(0o077)
    logging.basicConfig(level=logging.INFO)
    try:
        asyncio.run(main())
    except KeyboardInterrupt:
        pass
    except Exception as exc:
        log.error('%s', str(exc) if isinstance(exc, SafeError) else 'Inicialização falhou: ' + type(exc).__name__)
        raise SystemExit(1)
